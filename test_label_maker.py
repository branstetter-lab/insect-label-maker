import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
import openpyxl
import label_maker as l


class LabelTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        l.pdfmetrics.registerFont(l.TTFont('Arial', '/System/Library/Fonts/Supplemental/Arial.ttf'))
        l.pdfmetrics.registerFont(l.TTFont('LabelRegular', '/System/Library/Fonts/Supplemental/Arial.ttf'))
        l.pdfmetrics.registerFont(l.TTFont('LabelItalic', '/System/Library/Fonts/Supplemental/Arial Italic.ttf'))
        l.pdfmetrics.registerFont(l.TTFont('ArialBold', '/System/Library/Fonts/Supplemental/Arial Bold.ttf'))
        l.pdfmetrics.registerFont(l.TTFont('FieldHeading', '/System/Library/Fonts/Supplemental/Arial Bold.ttf'))

    def setUp(self):
        self.previous_main_size = l.MAIN_SIZE
        l.MAIN_SIZE = 3.5  # Legacy layout coverage; 4-point profile has its own test.

    def tearDown(self):
        l.MAIN_SIZE = self.previous_main_size

    def test_barcode_roundtrip(self):
        for code in ['BBSL1207527', 'ABC-000123']:
            self.assertEqual(l.zxingcpp.read_barcode(l.matrix(code)).text, code)

    def test_overflow_and_uncertainty(self):
        row = {'Catalog_Number': 'BBSL1207527', 'Field_Number': 'MGB4066',
               'Latitude': '41.91222', 'Longitude': '-111.67473', 'Uncertainty': '100'}
        lines = ['first', 'second', '[Latitude] [Longitude]±[Uncertainty]m', 'fourth', 'fifth', 'sixth']
        self.assertIsNone(l.layout(row, lines, l.MIN_WIDTH, True))
        self.assertIsNotNone(l.layout(row, lines, 90, True))
        self.assertNotIn('±', l.render_lines(row, lines, False)[2])
        self.assertIsNone(l.layout(row, ['extra'] * 7, 90, False))

    def test_fixed_size_six_line_layout(self):
        row = {'Catalog_Number': 'BBSL1207527', 'Field_Number': 'MGB4066',
               'Country': 'USA', 'StateProv': 'UT', 'County': 'Cache Co.',
               'LocalityLabel': 'Naomi Peak', 'Latitude': '41.91222',
               'Longitude': '-111.67473', 'Uncertainty': '100',
               'Verbatim_Elevation': '3000m', 'Verbatim_Date': '22Jul2021',
               'Collector_Label': 'M. Branstetter, T. Griswold',
               'MacroHabitat_Label': 'rocky peak', 'Microhabitat_Label': 'Phacelia sp.'}
        lines = l.layout(row, l.DEFAULT_LINES, l.MIN_WIDTH, False)
        self.assertIsNotNone(lines)
        self.assertEqual(len(lines), 6)
        self.assertEqual(lines[4], row['Collector_Label'])
        self.assertEqual(lines[5], 'rocky peak; Phacelia sp.')

    def test_default_barcode_size_and_center(self):
        img = l.matrix('BBSL1207527')
        bounds = img.point(lambda p: 255 - p).getbbox()
        dark_width = (bounds[2] - bounds[0]) * l.BAR / img.width
        self.assertAlmostEqual(dark_width / 72, .11 * 1.02)
        self.assertAlmostEqual(l.BAR_Y + l.BAR / 2, l.HEIGHT / 2)
        self.assertGreaterEqual(bounds[0], 2)
        self.assertGreaterEqual(img.width - bounds[2], 2)

    def test_four_point_profile_and_vertical_clearance(self):
        old_size = l.MAIN_SIZE
        try:
            l.MAIN_SIZE = 4
            row = {'Catalog_Number': 'BBSL1207604', 'Field_Number': 'MGB4039',
                   'Country': 'USA', 'StateProv': 'UT', 'County': 'Cache Co.',
                   'LocalityLabel': 'Tony Grove', 'Latitude': '41.89973',
                   'Longitude': '-111.64205', 'Uncertainty': '100',
                   'Verbatim_Elevation': '2510m', 'Verbatim_Date': '9Jul2021',
                   'Collector_Label': 'M.Branstetter, C.Meidt',
                   'MacroHabitat_Label': 'alpine meadow', 'Microhabitat_Label': 'pan traps'}
            lines = l.layout(row, l.DEFAULT_LINES, l.MIN_WIDTH, False)
            self.assertIsNotNone(lines)
            self.assertEqual(len(lines), 6)
            self.assertEqual([l.text_size(i) for i in range(6)], [4, 4, 4, 4, 4, 3.5])
            row['Microhabitat_Label'] = 'Geranium sp.'
            reduced = l.layout(row, l.DEFAULT_LINES, l.MIN_WIDTH, False)
            self.assertIsNotNone(reduced)
            self.assertEqual(reduced[-1], 'Geranium sp.')
            row['Microhabitat_Label'] = 'unreasonably long microhabitat ' * 20
            self.assertIsNone(l.layout(row, l.DEFAULT_LINES, l.MIN_WIDTH, False))
        finally:
            l.MAIN_SIZE = old_size

    def test_taxon_italics(self):
        for text, expected in [
            ('Phacelia sp.', ['Phacelia']),
            ('ex Phacelia tanacetifolia (Boraginaceae)', ['Phacelia', 'tanacetifolia']),
            ('Asteraceae sp.', []), ('Apiaceae sp.', []), ('Asteraceae', []),
            ('pan traps', []), ('air/ground', []),
            ('ex yellow Asteraceae', []),
        ]:
            self.assertEqual([text[a:b] for a, b in l.taxon_italic_ranges(text)], expected)
        lines = ['USA: UT', 'habitat; Phacelia sp.']
        mapped = l.label_italic_ranges(lines, 'Phacelia sp.')
        self.assertEqual(mapped[0], [])
        self.assertEqual([lines[1][a:b] for a, b in mapped[1]], ['Phacelia'])

    def test_adaptive_collector_and_flow(self):
        old_size, old_mode = l.MAIN_SIZE, l.LAYOUT_MODE
        try:
            l.MAIN_SIZE = 4
            row = {'Catalog_Number': 'BBSL1207527', 'Field_Number': 'MGB4066',
                   'Country': 'USA', 'StateProv': 'UT', 'County': 'Cache Co.',
                   'LocalityLabel': 'Naomi Peak', 'Latitude': '41.91222',
                   'Longitude': '-111.67473', 'Uncertainty': '100',
                   'Verbatim_Elevation': '3000m', 'Verbatim_Date': '22Jul2021',
                   'Collector_Label': 'M.Branstetter, T.Griswold',
                   'MacroHabitat_Label': 'rocky peak', 'Microhabitat_Label': 'Phacelia sp.'}
            l.LAYOUT_MODE = 'lines'
            lines = l.layout(row, l.DEFAULT_LINES, l.MIN_WIDTH, False)
            self.assertEqual(lines.sizes, [4, 4, 4, 4, 3.5, 3.5])
            self.assertEqual(lines[4], row['Collector_Label'])
            previous_bottom = None
            for index, (text, size, baseline) in enumerate(zip(lines, lines.sizes, lines.baselines)):
                above, below = l.ink_bounds(text, size)
                if previous_bottom is not None:
                    self.assertGreater(previous_bottom - baseline - above, 0)
                    self.assertAlmostEqual(lines.baselines[index - 1] - baseline, l.NORMAL_LEADING)
                previous_bottom = baseline - below
            short_row = dict(row, Collector_Label='M.Branstetter')
            short = l.layout(short_row, l.DEFAULT_LINES, l.MIN_WIDTH, False)
            self.assertEqual(short.sizes[4], 4)
            limit = l.MIN_WIDTH - 2 * l.PAD - l.SIDE - l.BAR - 1.7
            for text, size in zip(lines, lines.sizes):
                self.assertLessEqual(l.width(text, size), limit)
            l.LAYOUT_MODE = 'flow'
            flow = l.layout(row, l.DEFAULT_LINES, l.MIN_WIDTH, True)
            self.assertIsNotNone(flow)
            joined = ' '.join(flow)
            boundary_width = l.MIN_WIDTH - 2 * l.PAD - l.SIDE - l.BAR - 1.7
            for text, size in zip(flow, flow.sizes):
                self.assertLessEqual(l.width(text, size), boundary_width)
            for value in ['Naomi Peak', '±100m', 'M.Branstetter, T.Griswold', 'rocky peak, Phacelia sp.']:
                self.assertIn(value, joined)
            for text, size, baseline in zip(flow, flow.sizes, flow.baselines):
                self.assertLessEqual(l.width(text, size), l.available_at(l.MIN_WIDTH, baseline, text, size, row['Field_Number']))
                if 'Phacelia' in text:
                    self.assertEqual(size, 3.5)
        finally:
            l.MAIN_SIZE, l.LAYOUT_MODE = old_size, old_mode

    def test_condensed_coordinates_with_smaller_uncertainty(self):
        old_size, old_mode = l.MAIN_SIZE, l.LAYOUT_MODE
        try:
            l.MAIN_SIZE, l.LAYOUT_MODE = 4, 'lines'
            row = {'Latitude': '41.91222', 'Longitude': '-111.67473', 'Uncertainty': '500'}
            text = '41.91222 -111.67473±500m'
            runs = l.coordinate_runs(row, text, 4)
            self.assertEqual(runs[0][1:], (4, .61))
            self.assertEqual(runs[1], ('±500m', 3.5, .75))
            limit = l.MIN_WIDTH - 2 * l.PAD - l.SIDE - l.BAR - 1.7
            self.assertLessEqual(l.run_width(runs), limit)
            row['Uncertainty'] = ''
            self.assertEqual(len(l.coordinate_runs(row, '41.91222 -111.67473', 4)), 1)
        finally:
            l.MAIN_SIZE, l.LAYOUT_MODE = old_size, old_mode

    def test_width_limit_and_variable_height(self):
        old_height, old_bar_y = l.HEIGHT, l.BAR_Y
        try:
            for width in [.681, .75, 1, 0, -1, float('nan'), float('inf')]:
                with self.assertRaisesRegex(ValueError, 'width'):
                    l.configure_dimensions(width, .33)
            with self.assertRaisesRegex(ValueError, 'widening'):
                l.configure_dimensions(.68, .33, .7)
            for height in [.32, 11, float('nan')]:
                with self.assertRaisesRegex(ValueError, 'height'):
                    l.configure_dimensions(.68, height)
            l.configure_dimensions(.68, .4)
            self.assertAlmostEqual(l.HEIGHT, 28.8)
            self.assertAlmostEqual(l.BAR_Y + l.BAR / 2, l.HEIGHT / 2)
            l.configure_dimensions(.68, .33, .68)
            self.assertAlmostEqual(l.HEIGHT, 23.76)
        finally:
            l.HEIGHT, l.BAR_Y = old_height, old_bar_y

    def test_separate_habitat_templates_and_centering(self):
        main = l.separate_habitat_templates(l.DEFAULT_LINES)
        self.assertEqual(len(main), 5)
        self.assertFalse(any('Habitat_Label' in t for t in main))
        with self.assertRaisesRegex(ValueError, 'own template entries'):
            l.separate_habitat_templates(['[LocalityLabel], [MacroHabitat_Label]'])
        row = {'_excel_row': 2, 'Catalog_Number': 'DEMO000001',
               'MacroHabitat_Label': 'meadow', 'Microhabitat_Label': 'pan traps'}
        height = .15 * 72
        plan = l.habitat_label_plan(row, l.MIN_WIDTH, height)
        self.assertEqual(plan['font_points'], 4)
        self.assertEqual(plan['lines'], ['meadow', 'pan traps'])
        above, _ = l.ink_bounds('meadow', 4)
        _, below = l.ink_bounds('pan traps', 4)
        self.assertAlmostEqual(plan['baselines'][0] + above,
                               height - (plan['baselines'][1] - below))
        row['Microhabitat_Label'] = 'far too much habitat information ' * 10
        with self.assertRaisesRegex(ValueError, 'maximum is 3'):
            l.habitat_label_plan(row, l.MIN_WIDTH, height)

    def test_three_line_habitat_wrapping_and_italics(self):
        row = {'_excel_row': 2, 'Catalog_Number': 'DEMO000001',
               'MacroHabitat_Label': 'rocky slope beside alpine woodland',
               'Microhabitat_Label': 'Phacelia sp.'}
        plan = l.habitat_label_plan(row, .62 * 72, .2 * 72)
        self.assertEqual(len(plan['lines']), 3)
        self.assertEqual(plan['line_sources'], ['macrohabitat', 'macrohabitat', 'microhabitat'])
        self.assertEqual(' '.join(plan['lines']), row['MacroHabitat_Label'] + ' ' + row['Microhabitat_Label'])
        self.assertEqual(plan['lines'][2][slice(*plan['italic_ranges'][2][0])], 'Phacelia')
        self.assertEqual(plan['font_points'], 4)
        self.assertLess(plan['width_inches'], .68)
        with self.assertRaisesRegex(ValueError, 'maximum is 2'):
            l.habitat_label_plan(row, .62 * 72, .2 * 72, 2)
        with self.assertRaisesRegex(ValueError, 'too short'):
            l.habitat_label_plan(row, .62 * 72, .15 * 72)
        for text in plan['lines']:
            self.assertLessEqual(l.width(text, 4), .62 * 72 - 2 * l.PAD - .3)

    def test_separate_habitat_pairs_across_pages(self):
        from pypdf import PdfReader
        script = Path(l.__file__).resolve()
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = openpyxl.load_workbook(script.parent / 'examples/example_specimens.xlsx')
            sheet = book.active
            sample = list(next(sheet.iter_rows(min_row=2, max_row=2, values_only=True)))
            sheet.delete_rows(2, sheet.max_row)
            for i in range(221):
                row = list(sample)
                row[0] = f'DEMO{i:06d}'
                sheet.append(row)
            source = root / 'specimens.xlsx'
            output = root / 'pairs.pdf'
            book.save(source)
            result = subprocess.run([sys.executable, str(script), str(source), '--separate-habitat',
                                     '--output', str(output)], capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(output.with_suffix('.json').read_text())
            self.assertEqual(report['physical_labels'], 442)
            self.assertEqual(report['barcode_rendering'], 'raster')
            self.assertEqual(report['barcode_raster_dpi'], 1200)
            self.assertEqual(report['barcode_scale'], 1.02)
            self.assertEqual(report['coordinate_horizontal_scale'], .61)
            self.assertEqual(report['label_rendering'], 'vector')
            self.assertEqual(report['habitat_font_points'], 4)
            self.assertTrue(report['field_font_path'].endswith('Verdana.ttf'))
            self.assertEqual(report['field_font_points'], 4)
            self.assertEqual(len(PdfReader(output).pages), 2)
            capacity = int((792 - 36 + l.PAGE_GAP) // ((.33 + .02 + .20) * 72 + l.PAGE_GAP)) * 11
            self.assertEqual(report['records'][capacity - 1]['habitat_label']['page'], 1)
            self.assertEqual(report['records'][capacity]['habitat_label']['page'], 2)
            self.assertEqual(report['habitat_label_width_inches'], .62)
            for record in report['records']:
                self.assertEqual(len(record['lines']), 5)
                plan = record['habitat_label']
                self.assertGreaterEqual(plan['y_points'], 18)
                self.assertEqual(plan['font_points'], 4)
                self.assertEqual(plan['lines'], ['meadow', 'pan traps'])

    def test_field_block_headings_and_column_continuation(self):
        rows = [{'Field_Number': code} for code in ['EVT001', 'EVT001', 'EVT002', 'EVT001']]
        positions = l.sheet_positions(rows, l.MIN_WIDTH, l.HEIGHT, True)
        self.assertEqual([bool(p['heading']) for p in positions], [True, False, True, True])
        self.assertEqual([p['heading']['field_number'] for p in positions if p['heading']],
                         ['EVT001', 'EVT002', 'EVT001'])
        for p in positions:
            if p['heading']:
                self.assertGreater(p['heading']['baseline'], p['y_points'] + l.HEIGHT)
        long_block = l.sheet_positions([{'Field_Number': 'EVT001'}] * 400, l.MIN_WIDTH, l.HEIGHT, True)
        self.assertGreater(long_block[-1]['page'], 1)
        self.assertTrue(any(p['heading'] and p['heading']['continued'] for p in long_block))
        for previous, p in zip(long_block, long_block[1:]):
            if p['page'] != previous['page'] or p['x_points'] != previous['x_points']:
                self.assertIsNotNone(p['heading'])
                self.assertTrue(p['heading']['continued'])
        self.assertTrue(all(p['heading'] is None for p in l.sheet_positions(rows, l.MIN_WIDTH, l.HEIGHT, False)))

    def test_field_number_fits_beside_text(self):
        right_edge = l.MIN_WIDTH - l.PAD - l.SIDE
        text_edge = right_edge - l.BAR - 1.7
        code = 'MGB4066'
        self.assertGreater(right_edge - l.field_width(code) - l.FIELD_LEFT_SHIFT, text_edge)
        self.assertEqual(l.FIELD_SIZE, 4)

    def test_all_verdana_output(self):
        from pypdf import PdfReader
        script = Path(l.__file__).resolve()
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / 'verdana.pdf'
            result = subprocess.run([sys.executable, str(script), str(script.parent / 'examples/example_specimens.xlsx'),
                                     '--font-family', 'verdana', '--separate-habitat', '-o', str(output)],
                                    capture_output=True, text=True)
            self.assertEqual(result.returncode, 0, result.stderr)
            report = json.loads(output.with_suffix('.json').read_text())
            self.assertEqual(report['font'], 'Verdana')
            self.assertEqual(report['width_inches'], .68)
            self.assertEqual(report['coordinate_horizontal_scale'], .49)
            self.assertEqual(report['uncertainty_omitted_count'], 0)
            used_fonts = []
            def visitor(text, cm, tm, font, size):
                if text.strip():
                    used_fonts.append(str(font.get('/BaseFont', '')))
            for page in PdfReader(output).pages:
                page.extract_text(visitor_text=visitor)
            self.assertTrue(used_fonts)
            self.assertTrue(all('Verdana' in font for font in used_fonts))
            self.assertTrue(any('Italic' in font for font in used_fonts))
            self.assertTrue(any('Bold' in font for font in used_fonts))
            from pypdf.generic import ContentStream
            charspace, stack = 0, []
            checked = 0
            for page in PdfReader(output).pages:
                for operands, operator in ContentStream(page.get_contents(), page.pdf).operations:
                    if operator == b'q':
                        stack.append(charspace)
                    elif operator == b'Q':
                        charspace = stack.pop()
                    elif operator == b'Tc':
                        charspace = float(operands[0])
                    elif operator == b'Tj' and str(operands[0]).strip('- '):
                        self.assertEqual(charspace, 0)
                        checked += 1
            self.assertGreater(checked, 0)

    def test_heading_dash_segments_fill_width(self):
        plan = l.heading_segments('MGB4066', l.MIN_WIDTH)
        used = 2 * plan['side_width'] + 2 * plan['gap'] + plan['number_width']
        self.assertAlmostEqual(used, l.MIN_WIDTH - 2 * l.PAD)
        dash_width = l.pdfmetrics.stringWidth('-', 'FieldHeading', plan['dash_size'])
        side = plan['dash_count'] * dash_width + (plan['dash_count'] - 1) * plan['dash_spacing']
        self.assertAlmostEqual(side, plan['side_width'])
        self.assertLess(plan['dash_size'], plan['size'])

    def test_habitat_clearance_from_field_number(self):
        self.assertEqual(l.FIELD_LEFT_SHIFT, 0)
        self.assertEqual(l.HABITAT_FIELD_GAP, 1.5)
        for code in ['MGB4066', 'EVT001']:
            available = l.habitat_available(l.MIN_WIDTH, code)
            field_left = l.MIN_WIDTH - l.PAD - l.SIDE - l.field_width(code) - l.FIELD_LEFT_SHIFT
            self.assertGreaterEqual(field_left - (l.PAD + available), l.HABITAT_FIELD_GAP - 1e-9)
            self.assertLessEqual(available, l.line_available(5, l.MIN_WIDTH))
        self.assertEqual(l.habitat_available(l.MIN_WIDTH, ''), l.line_available(5, l.MIN_WIDTH))

    def test_duplicate_ids_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'input.xlsx'
            wb = openpyxl.Workbook()
            wb.active.append(['Catalog_Number'])
            wb.active.append(['ID1'])
            wb.active.append(['ID1'])
            wb.save(p)
            with self.assertRaisesRegex(ValueError, 'Duplicate Catalog_Number'):
                l.read_rows(p)

if __name__ == '__main__':
    unittest.main()
